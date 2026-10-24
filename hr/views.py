import io
import ipaddress
import math
import re
from datetime import date, timedelta

from django.contrib import messages
from django.contrib.auth import get_user_model
from django.contrib.auth.decorators import (
    login_required,
    permission_required,
    user_passes_test,
)
from django.contrib.auth.models import Group, Permission
from django.db import transaction
from django.db.models import Q
from django.http import HttpResponse
from django.shortcuts import render, redirect, get_object_or_404
from django.utils import timezone

import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from reportlab.lib.pagesizes import letter, landscape
from reportlab.lib import colors
from reportlab.platypus import SimpleDocTemplate, Paragraph, Table, TableStyle
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle

from .models import (
    Department,
    Employee,
    Attendance,
    LeaveRequest,
    HospitalSetting,
)
from .forms import EmployeeForm, LeaveRequestForm


def calculate_haversine_distance(lat1, lon1, lat2, lon2):
    """
    Calculate distance in meters between two coordinates on Earth using Haversine formula.
    """
    R = 6371000  # meters
    phi1 = math.radians(lat1)
    phi2 = math.radians(lat2)
    delta_phi = math.radians(lat2 - lat1)
    delta_lambda = math.radians(lon2 - lon1)

    a = (
        math.sin(delta_phi / 2.0) ** 2
        + math.cos(phi1) * math.cos(phi2) * (math.sin(delta_lambda / 2.0) ** 2)
    )
    c = 2 * math.atan2(math.sqrt(a), math.sqrt(1 - a))
    return R * c


def get_client_ip(request):
    """
    Extract the client's IP address, handling proxies and forwarded headers.
    """
    x_forwarded_for = request.META.get("HTTP_X_FORWARDED_FOR")
    if x_forwarded_for:
        ip = x_forwarded_for.split(",")[0].strip()
    else:
        ip = request.META.get("REMOTE_ADDR", "").strip()
    return ip


def is_ip_allowed(client_ip, allowed_ips_text):
    """
    Check if client IP matches allowed Wi-Fi IP addresses, wildcards (e.g. 192.168.1.*),
    or CIDR network ranges (e.g. 192.168.1.0/24).
    """
    if not client_ip or not allowed_ips_text:
        return False

    client_ip = client_ip.strip()
    client_ip_clean = "127.0.0.1" if client_ip in ("::1", "127.0.0.1", "localhost") else client_ip

    tokens = [t.strip() for t in re.split(r"[\r\n,]+", allowed_ips_text) if t.strip()]

    for token in tokens:
        # Match localhost aliases
        if token.lower() in ("localhost", "127.0.0.1", "::1") and client_ip_clean == "127.0.0.1":
            return True

        # Exact IP match
        if client_ip == token or client_ip_clean == token:
            return True

        # Wildcard prefix match (e.g., 192.168.1.* or 10.0.*)
        if "*" in token:
            prefix = token.split("*")[0]
            if client_ip.startswith(prefix) or client_ip_clean.startswith(prefix):
                return True

        # CIDR notation match (e.g., 192.168.1.0/24)
        if "/" in token:
            try:
                network = ipaddress.ip_network(token, strict=False)
                if ipaddress.ip_address(client_ip_clean) in network:
                    return True
            except ValueError:
                pass

    return False



User = get_user_model()


# =====================================================
# HELPER: CREATE OR UPDATE AN EMPLOYEE LOGIN ACCOUNT
# =====================================================
def create_employee_user(employee, username, password=None):
    username = username.strip()

    if not username:
        raise ValueError("Username is required.")

    existing_users = User.objects.filter(
        username__iexact=username
    )

    if employee.user_id:
        existing_users = existing_users.exclude(
            pk=employee.user_id
        )

    if existing_users.exists():
        raise ValueError(
            "This username already exists. Please choose another."
        )

    if employee.user_id:
        user = employee.user
        user.username = username
        user.email = employee.email
        user.first_name = employee.name

        if password:
            user.set_password(password)

        user.save()
        return user

    if not password:
        raise ValueError(
            "A password is required to create a new login account."
        )

    # Create the Django account first.
    user = User.objects.create_user(
        username=username,
        email=employee.email,
        password=password,
        first_name=employee.name,
    )

    # Attach the user to the employee.
    employee.user = user

    # If the employee has not been inserted yet, save the complete
    # employee record. Otherwise, update its existing record.
    if employee.pk:
        employee.save(update_fields=["user"])
    else:
        employee.save()

    return user


# =====================================================
# 1. DASHBOARD
# =====================================================

@login_required
@permission_required("hr.view_dashboard", raise_exception=True)
def dashboard(request):
    today = timezone.localdate()

    can_view_leaves = request.user.has_perm(
        "hr.view_leave_pages"
    )

    context = {
        "today": today,

        "total_employees": Employee.objects.count(),

        "active_employees": Employee.objects.filter(
            status="ACTIVE"
        ).count(),

        "total_departments": Department.objects.count(),

        "present_today": Attendance.objects.filter(
            date=today,
            check_in__isnull=False,
        ).count(),

        "pending_leaves": (
            LeaveRequest.objects.filter(
                status="PENDING"
            ).count()
            if can_view_leaves
            else None
        ),

        "recent_leaves": (
            LeaveRequest.objects.select_related(
                "employee"
            ).order_by("-start_date")[:5]
            if can_view_leaves
            else []
        ),
    }

    return render(
        request,
        "hr/dashboard.html",
        context,
    )


# =====================================================
# 2. EMPLOYEE LIST
# =====================================================

@login_required
@permission_required("hr.view_employee_pages", raise_exception=True)
def employee_list(request):
    employees = Employee.objects.select_related(
        "department",
        "user",
    ).all().order_by("name")

    departments = Department.objects.all()

    query = request.GET.get("q", "").strip()
    department = request.GET.get("department", "")
    status = request.GET.get("status", "")

    if query:
        employees = employees.filter(
            Q(name__icontains=query)
            | Q(employee_id__icontains=query)
            | Q(email__icontains=query)
        )

    if department:
        employees = employees.filter(
            department_id=department
        )

    if status in ("ACTIVE", "INACTIVE", "ON_LEAVE"):
        employees = employees.filter(
            status=status
        )

    return render(
        request,
        "hr/employee_list.html",
        {
            "employees": employees,
            "departments": departments,
            "query": query,
            "selected_department": department,
            "selected_status": status,
        },
    )


# =====================================================
# 3. ADD EMPLOYEE
# =====================================================

@login_required
@permission_required("hr.manage_employee_pages", raise_exception=True)
def employee_create(request):
    if request.method == "POST":
        form = EmployeeForm(request.POST)

        if form.is_valid():
            try:
                with transaction.atomic():
                    employee = form.save(commit=False)

                    username = form.cleaned_data["username"]
                    password = form.cleaned_data["password"]
                    selected_role = form.cleaned_data["role"]

                    # Create the login account with the username
                    # and password entered in the form.
                    user = create_employee_user(
                        employee=employee,
                        username=username,
                        password=password,
                    )

                    employee.user = user
                    employee.save()

                    # Only superusers can assign employee roles.
                    if request.user.is_superuser:
                        user.groups.set([selected_role])

                messages.success(
                    request,
                    f"Employee '{employee.name}' and login account "
                    "created successfully.",
                )

                return redirect(
                    "employee_profile",
                    pk=employee.pk,
                )

            except Exception as error:
                messages.error(
                    request,
                    f"Unable to create employee: {error}",
                )

    else:
        form = EmployeeForm()

    return render(
        request,
        "hr/employee_form.html",
        {
            "form": form,
            "page_title": "Add Employee",
            "title": "Add Employee",
        },
    )


# =====================================================
# 4. EDIT EMPLOYEE
# =====================================================

@login_required
@permission_required("hr.manage_employee_pages", raise_exception=True)
def employee_edit(request, pk):
    employee = get_object_or_404(
        Employee.objects.select_related(
            "department",
            "user",
        ),
        pk=pk,
    )

    if request.method == "POST":
        form = EmployeeForm(
            request.POST,
            instance=employee,
        )

        if form.is_valid():
            try:
                with transaction.atomic():
                    employee = form.save(commit=False)

                    username = form.cleaned_data["username"]
                    password = form.cleaned_data.get("password")
                    selected_role = form.cleaned_data["role"]

                    # Update the existing account or create one
                    # if the employee does not have a linked user.
                    user = create_employee_user(
                        employee=employee,
                        username=username,
                        password=password,
                    )

                    employee.user = user
                    employee.save()

                    if request.user.is_superuser:
                        user.groups.set([selected_role])

                messages.success(
                    request,
                    "Employee details updated successfully.",
                )

                return redirect(
                    "employee_profile",
                    pk=employee.pk,
                )

            except Exception as error:
                messages.error(
                    request,
                    f"Unable to update employee: {error}",
                )

    else:
        form = EmployeeForm(
            instance=employee,
        )

    return render(
        request,
        "hr/employee_form.html",
        {
            "form": form,
            "page_title": "Edit Employee",
            "title": "Edit Employee",
        },
    )


# =====================================================
# 5. EMPLOYEE PROFILE
# =====================================================

@login_required
@permission_required("hr.view_employee_pages", raise_exception=True)
def employee_profile(request, pk):
    employee = get_object_or_404(
        Employee.objects.select_related(
            "department",
            "user",
        ),
        pk=pk,
    )

    attendance_records = Attendance.objects.filter(
        employee=employee
    )

    can_view_leaves = request.user.has_perm(
        "hr.view_leave_pages"
    )

    context = {
        "employee": employee,

        "days_worked": attendance_records.filter(
            check_in__isnull=False
        ).values("date").distinct().count(),

        "attendance_count": attendance_records.count(),

        "leave_requests": (
            LeaveRequest.objects.filter(
                employee=employee
            ).order_by("-start_date")[:5]
            if can_view_leaves
            else []
        ),
    }

    return render(
        request,
        "hr/employee_profile.html",
        context,
    )


def export_attendance_excel(attendance_records, selected_date, department_name):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Attendance Report"

    # Title block
    ws.merge_cells("A1:G1")
    title_cell = ws["A1"]
    title_cell.value = "Hospital HRMS - Daily Attendance Report"
    title_cell.font = Font(name="Calibri", size=15, bold=True, color="1E3A8A")
    title_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 28

    # Subtitle / Filter info
    ws.merge_cells("A2:G2")
    sub_cell = ws["A2"]
    sub_cell.value = (
        f"Date: {selected_date.strftime('%d-%b-%Y')}  |  "
        f"Department: {department_name}  |  "
        f"Generated on: {timezone.now().strftime('%d-%b-%Y %H:%M')}"
    )
    sub_cell.font = Font(name="Calibri", size=10, italic=True, color="475569")
    sub_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[2].height = 20

    # Header Row
    headers = [
        "Employee ID",
        "Employee Name",
        "Department",
        "Shift",
        "Check In",
        "Check Out",
        "Hours Worked",
    ]
    ws.row_dimensions[4].height = 24
    header_fill = PatternFill(
        start_color="1E3A8A", end_color="1E3A8A", fill_type="solid"
    )
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    thin_border = Border(
        left=Side(style="thin", color="CBD5E1"),
        right=Side(style="thin", color="CBD5E1"),
        top=Side(style="thin", color="CBD5E1"),
        bottom=Side(style="thin", color="CBD5E1"),
    )

    for col_idx, header in enumerate(headers, start=1):
        cell = ws.cell(row=4, column=col_idx, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = thin_border

    # Data Rows
    row_idx = 5
    alt_fill = PatternFill(
        start_color="F8FAFC", end_color="F8FAFC", fill_type="solid"
    )

    for record in attendance_records:
        ws.row_dimensions[row_idx].height = 20
        check_in_str = (
            record.check_in.strftime("%H:%M") if record.check_in else "—"
        )
        check_out_str = (
            record.check_out.strftime("%H:%M") if record.check_out else "—"
        )
        shift_str = record.shift.name if record.shift else "—"
        dept_str = (
            record.employee.department.name
            if record.employee and record.employee.department
            else "—"
        )

        row_data = [
            record.employee.employee_id if record.employee else "—",
            record.employee.name if record.employee else "—",
            dept_str,
            shift_str,
            check_in_str,
            check_out_str,
            record.hours_worked,
        ]

        for col_idx, val in enumerate(row_data, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=val)
            cell.font = Font(name="Calibri", size=10)
            cell.border = thin_border
            if row_idx % 2 == 0:
                cell.fill = alt_fill
            if col_idx in (1, 4, 5, 6, 7):
                cell.alignment = Alignment(
                    horizontal="center", vertical="center"
                )
            else:
                cell.alignment = Alignment(
                    horizontal="left", vertical="center"
                )

        row_idx += 1

    if not attendance_records.exists():
        ws.merge_cells(
            start_row=row_idx,
            start_column=1,
            end_row=row_idx,
            end_column=7,
        )
        empty_cell = ws.cell(
            row=row_idx,
            column=1,
            value="No attendance records found for this date/filter.",
        )
        empty_cell.font = Font(
            name="Calibri", size=10, italic=True, color="64748B"
        )
        empty_cell.alignment = Alignment(
            horizontal="center", vertical="center"
        )
        ws.row_dimensions[row_idx].height = 24

    col_widths = [16, 26, 22, 16, 14, 14, 16]
    for col_idx, width in enumerate(col_widths, start=1):
        col_letter = openpyxl.utils.get_column_letter(col_idx)
        ws.column_dimensions[col_letter].width = width

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    filename = f"Attendance_Report_{selected_date.strftime('%Y-%m-%d')}.xlsx"
    response = HttpResponse(
        buffer.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    return response


def export_attendance_pdf(attendance_records, selected_date, department_name):
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=landscape(letter),
        leftMargin=36,
        rightMargin=36,
        topMargin=36,
        bottomMargin=36,
    )

    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        "PDFTitle",
        parent=styles["Heading1"],
        fontName="Helvetica-Bold",
        fontSize=18,
        textColor=colors.HexColor("#1e3a8a"),
        alignment=1,
        spaceAfter=4,
    )
    sub_style = ParagraphStyle(
        "PDFSub",
        parent=styles["Normal"],
        fontName="Helvetica",
        fontSize=10,
        textColor=colors.HexColor("#475569"),
        alignment=1,
        spaceAfter=15,
    )
    cell_style = ParagraphStyle(
        "PDFCell",
        parent=styles["Normal"],
        fontName="Helvetica",
        fontSize=9,
        textColor=colors.HexColor("#1e293b"),
        alignment=0,
    )
    cell_center = ParagraphStyle(
        "PDFCellCenter",
        parent=cell_style,
        alignment=1,
    )
    header_style = ParagraphStyle(
        "PDFHeader",
        parent=styles["Normal"],
        fontName="Helvetica-Bold",
        fontSize=9,
        textColor=colors.white,
        alignment=1,
    )

    elements = []
    elements.append(
        Paragraph("Hospital HRMS - Daily Attendance Report", title_style)
    )
    elements.append(
        Paragraph(
            f"Date: <b>{selected_date.strftime('%d %B %Y')}</b> &nbsp;&nbsp;|&nbsp;&nbsp; "
            f"Department: <b>{department_name}</b> &nbsp;&nbsp;|&nbsp;&nbsp; "
            f"Total Listed: <b>{attendance_records.count()}</b>",
            sub_style,
        )
    )

    table_data = [
        [
            Paragraph("Employee ID", header_style),
            Paragraph("Name", header_style),
            Paragraph("Department", header_style),
            Paragraph("Shift", header_style),
            Paragraph("Check In", header_style),
            Paragraph("Check Out", header_style),
            Paragraph("Hours", header_style),
        ]
    ]

    for record in attendance_records:
        check_in_str = (
            record.check_in.strftime("%H:%M") if record.check_in else "—"
        )
        check_out_str = (
            record.check_out.strftime("%H:%M") if record.check_out else "—"
        )
        shift_str = record.shift.name if record.shift else "—"
        dept_str = (
            record.employee.department.name
            if record.employee and record.employee.department
            else "—"
        )

        table_data.append([
            Paragraph(
                record.employee.employee_id if record.employee else "—",
                cell_center,
            ),
            Paragraph(
                record.employee.name if record.employee else "—", cell_style
            ),
            Paragraph(dept_str, cell_style),
            Paragraph(shift_str, cell_center),
            Paragraph(check_in_str, cell_center),
            Paragraph(check_out_str, cell_center),
            Paragraph(record.hours_worked, cell_center),
        ])

    if len(table_data) == 1:
        empty_style = ParagraphStyle(
            "PDFEmpty",
            parent=cell_center,
            fontName="Helvetica-Oblique",
            textColor=colors.HexColor("#64748b"),
        )
        table_data.append([
            Paragraph(
                "No attendance records found for this filter.", empty_style
            ),
            "",
            "",
            "",
            "",
            "",
            "",
        ])

    col_widths = [85, 150, 125, 90, 85, 85, 100]
    t = Table(table_data, colWidths=col_widths, repeatRows=1)

    t_style = [
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1e3a8a")),
        ("ALIGN", (0, 0), (-1, -1), "CENTER"),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("BOTTOMPADDING", (0, 0), (-1, 0), 8),
        ("TOPPADDING", (0, 0), (-1, 0), 8),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#cbd5e1")),
        (
            "ROWBACKGROUNDS",
            (0, 1),
            (-1, -1),
            [colors.white, colors.HexColor("#f8fafc")],
        ),
        ("TOPPADDING", (0, 1), (-1, -1), 6),
        ("BOTTOMPADDING", (0, 1), (-1, -1), 6),
    ]
    if len(table_data) == 2 and not attendance_records.exists():
        t_style.append(("SPAN", (0, 1), (6, 1)))

    t.setStyle(TableStyle(t_style))
    elements.append(t)

    doc.build(elements)
    buffer.seek(0)

    filename = f"Attendance_Report_{selected_date.strftime('%Y-%m-%d')}.pdf"
    response = HttpResponse(buffer.getvalue(), content_type="application/pdf")
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    return response


def is_hr_or_admin(user):
    if not user.is_authenticated:
        return False
    if user.is_superuser or user.is_staff:
        return True
    return user.groups.filter(
        name__iregex=r"^(superadmin|admin|hr|hr\s*manager|hospital\s*admin)$"
    ).exists()


def export_personal_attendance_excel(attendance_records, employee, selected_month_label):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "My Attendance"

    # Title block
    ws.merge_cells("A1:F1")
    title_cell = ws["A1"]
    title_cell.value = "Hospital HRMS - Personal Attendance Record"
    title_cell.font = Font(name="Calibri", size=15, bold=True, color="1E3A8A")
    title_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 28

    # Subtitle / Employee info
    ws.merge_cells("A2:F2")
    sub_cell = ws["A2"]
    dept_name = employee.department.name if employee.department else "—"
    sub_cell.value = (
        f"Employee: {employee.name} ({employee.employee_id})  |  "
        f"Department: {dept_name}  |  "
        f"Designation: {employee.designation}  |  "
        f"Month: {selected_month_label}"
    )
    sub_cell.font = Font(name="Calibri", size=10, italic=True, color="475569")
    sub_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[2].height = 20

    # Header Row
    headers = [
        "Date",
        "Day",
        "Shift",
        "Check In",
        "Check Out",
        "Hours Worked",
    ]
    ws.row_dimensions[4].height = 24
    header_fill = PatternFill(
        start_color="1E3A8A", end_color="1E3A8A", fill_type="solid"
    )
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    thin_border = Border(
        left=Side(style="thin", color="CBD5E1"),
        right=Side(style="thin", color="CBD5E1"),
        top=Side(style="thin", color="CBD5E1"),
        bottom=Side(style="thin", color="CBD5E1"),
    )

    for col_idx, header in enumerate(headers, start=1):
        cell = ws.cell(row=4, column=col_idx, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = thin_border

    # Data Rows
    row_idx = 5
    alt_fill = PatternFill(
        start_color="F8FAFC", end_color="F8FAFC", fill_type="solid"
    )

    for record in attendance_records:
        ws.row_dimensions[row_idx].height = 20
        date_str = record.date.strftime("%d-%b-%Y")
        day_str = record.date.strftime("%A")
        check_in_str = (
            record.check_in.strftime("%H:%M") if record.check_in else "—"
        )
        check_out_str = (
            record.check_out.strftime("%H:%M") if record.check_out else "—"
        )
        shift_str = record.shift.name if record.shift else "—"

        row_data = [
            date_str,
            day_str,
            shift_str,
            check_in_str,
            check_out_str,
            record.hours_worked,
        ]

        for col_idx, val in enumerate(row_data, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=val)
            cell.font = Font(name="Calibri", size=10)
            cell.border = thin_border
            cell.alignment = Alignment(
                horizontal="center", vertical="center"
            )
            if row_idx % 2 == 0:
                cell.fill = alt_fill

        row_idx += 1

    if not attendance_records.exists():
        ws.merge_cells(
            start_row=row_idx,
            start_column=1,
            end_row=row_idx,
            end_column=6,
        )
        empty_cell = ws.cell(
            row=row_idx,
            column=1,
            value="No attendance records found for this period.",
        )
        empty_cell.font = Font(
            name="Calibri", size=10, italic=True, color="64748B"
        )
        empty_cell.alignment = Alignment(
            horizontal="center", vertical="center"
        )
        ws.row_dimensions[row_idx].height = 24

    col_widths = [16, 16, 18, 14, 14, 16]
    for col_idx, width in enumerate(col_widths, start=1):
        col_letter = openpyxl.utils.get_column_letter(col_idx)
        ws.column_dimensions[col_letter].width = width

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    filename = f"My_Attendance_{employee.employee_id}.xlsx"
    response = HttpResponse(
        buffer.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    return response


def export_personal_attendance_pdf(attendance_records, employee, selected_month_label):
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=letter,
        leftMargin=36,
        rightMargin=36,
        topMargin=36,
        bottomMargin=36,
    )

    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        "PDFTitle",
        parent=styles["Heading1"],
        fontName="Helvetica-Bold",
        fontSize=16,
        textColor=colors.HexColor("#1e3a8a"),
        alignment=1,
        spaceAfter=4,
    )
    sub_style = ParagraphStyle(
        "PDFSub",
        parent=styles["Normal"],
        fontName="Helvetica",
        fontSize=10,
        textColor=colors.HexColor("#475569"),
        alignment=1,
        spaceAfter=15,
    )
    cell_style = ParagraphStyle(
        "PDFCell",
        parent=styles["Normal"],
        fontName="Helvetica",
        fontSize=9,
        textColor=colors.HexColor("#1e293b"),
        alignment=1,
    )
    header_style = ParagraphStyle(
        "PDFHeader",
        parent=styles["Normal"],
        fontName="Helvetica-Bold",
        fontSize=9,
        textColor=colors.white,
        alignment=1,
    )

    elements = []
    elements.append(
        Paragraph("Hospital HRMS - Personal Attendance Record", title_style)
    )
    dept_name = employee.department.name if employee.department else "—"
    elements.append(
        Paragraph(
            f"Employee: <b>{employee.name}</b> ({employee.employee_id}) &nbsp;|&nbsp; "
            f"Dept: <b>{dept_name}</b> &nbsp;|&nbsp; "
            f"Designation: <b>{employee.designation}</b> &nbsp;|&nbsp; "
            f"Month: <b>{selected_month_label}</b>",
            sub_style,
        )
    )

    table_data = [
        [
            Paragraph("Date", header_style),
            Paragraph("Day", header_style),
            Paragraph("Shift", header_style),
            Paragraph("Check In", header_style),
            Paragraph("Check Out", header_style),
            Paragraph("Hours", header_style),
        ]
    ]

    for record in attendance_records:
        check_in_str = (
            record.check_in.strftime("%H:%M") if record.check_in else "—"
        )
        check_out_str = (
            record.check_out.strftime("%H:%M") if record.check_out else "—"
        )
        shift_str = record.shift.name if record.shift else "—"

        table_data.append([
            Paragraph(record.date.strftime("%d-%b-%Y"), cell_style),
            Paragraph(record.date.strftime("%A"), cell_style),
            Paragraph(shift_str, cell_style),
            Paragraph(check_in_str, cell_style),
            Paragraph(check_out_str, cell_style),
            Paragraph(record.hours_worked, cell_style),
        ])

    if len(table_data) == 1:
        empty_style = ParagraphStyle(
            "PDFEmpty",
            parent=cell_style,
            fontName="Helvetica-Oblique",
            textColor=colors.HexColor("#64748b"),
        )
        table_data.append([
            Paragraph(
                "No attendance records found for this period.", empty_style
            ),
            "",
            "",
            "",
            "",
            "",
        ])

    col_widths = [90, 85, 110, 85, 85, 85]
    t = Table(table_data, colWidths=col_widths, repeatRows=1)

    t_style = [
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1e3a8a")),
        ("ALIGN", (0, 0), (-1, -1), "CENTER"),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("BOTTOMPADDING", (0, 0), (-1, 0), 8),
        ("TOPPADDING", (0, 0), (-1, 0), 8),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#cbd5e1")),
        (
            "ROWBACKGROUNDS",
            (0, 1),
            (-1, -1),
            [colors.white, colors.HexColor("#f8fafc")],
        ),
        ("TOPPADDING", (0, 1), (-1, -1), 6),
        ("BOTTOMPADDING", (0, 1), (-1, -1), 6),
    ]
    if len(table_data) == 2 and not attendance_records.exists():
        t_style.append(("SPAN", (0, 1), (5, 1)))

    t.setStyle(TableStyle(t_style))
    elements.append(t)

    doc.build(elements)
    buffer.seek(0)

    filename = f"My_Attendance_{employee.employee_id}.pdf"
    response = HttpResponse(buffer.getvalue(), content_type="application/pdf")
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    return response


# =====================================================
# 6. ATTENDANCE REPORT & EMPLOYEE ATTENDANCE
# =====================================================

@login_required
def attendance_report(request):
    user_is_admin = is_hr_or_admin(request.user)
    requested_view = request.GET.get("view", "")
    hospital_setting = HospitalSetting.get_settings()
    client_ip = get_client_ip(request)
    is_on_hospital_wifi = is_ip_allowed(client_ip, hospital_setting.allowed_wifi_ips)

    # ---------------------------------------------------------
    # ADMIN ACTIONS: Location & Wi-Fi Management
    # ---------------------------------------------------------
    if request.method == "POST" and user_is_admin:
        action = request.POST.get("action")

        # 1. Update Hospital GPS Location
        if action == "set_hospital_location":
            try:
                new_lat = float(request.POST.get("latitude"))
                new_lon = float(request.POST.get("longitude"))
                hospital_setting.latitude = new_lat
                hospital_setting.longitude = new_lon
                hospital_setting.save(
                    update_fields=["latitude", "longitude"]
                )
                messages.success(
                    request,
                    f"Hospital GPS location successfully set to ({new_lat:.6f}, {new_lon:.6f}).",
                )
            except (ValueError, TypeError):
                messages.error(
                    request, "Invalid GPS coordinates provided."
                )
            return redirect("attendance_report")

        # 2. Register Current Device's Wi-Fi as Authorized Hospital Wi-Fi
        elif action == "register_hospital_wifi":
            current_ips = [
                ip.strip()
                for ip in re.split(r"[\r\n,]+", hospital_setting.allowed_wifi_ips)
                if ip.strip()
            ]

            to_add = [client_ip]
            # If standard private IPv4, add wildcard subnet as well (e.g. 192.168.1.*)
            parts = client_ip.split(".")
            if len(parts) == 4 and parts[0] in ("192", "10", "172"):
                subnet = f"{parts[0]}.{parts[1]}.{parts[2]}.*"
                if subnet not in current_ips:
                    to_add.append(subnet)

            for ip_cand in to_add:
                if ip_cand not in current_ips:
                    current_ips.append(ip_cand)

            hospital_setting.allowed_wifi_ips = ", ".join(current_ips)
            wifi_label = request.POST.get("wifi_name", "").strip()
            if wifi_label:
                hospital_setting.wifi_name = wifi_label
            hospital_setting.save()

            messages.success(
                request,
                f"Current network (IP: {client_ip}) successfully registered as authorized '{hospital_setting.wifi_name}'. Staff on this network can now mark attendance.",
            )
            return redirect("attendance_report")

        # 3. Change Verification Mode (Wi-Fi Only, GPS Only, Either, Both)
        elif action == "update_verification_settings":
            new_mode = request.POST.get("verification_mode")
            new_wifi_name = request.POST.get("wifi_name", "").strip()
            new_allowed_ips = request.POST.get("allowed_wifi_ips", "").strip()
            new_radius = request.POST.get("radius_meters")

            if new_mode in dict(HospitalSetting.VerificationMode.choices):
                hospital_setting.verification_mode = new_mode
            if new_wifi_name:
                hospital_setting.wifi_name = new_wifi_name
            if new_allowed_ips:
                hospital_setting.allowed_wifi_ips = new_allowed_ips
            if new_radius:
                try:
                    hospital_setting.radius_meters = int(new_radius)
                except ValueError:
                    pass

            hospital_setting.save()
            messages.success(request, f"Attendance verification mode updated to: {hospital_setting.get_verification_mode_display()}.")
            return redirect("attendance_report")

    # Regular employees ALWAYS see their own personal attendance.
    # HR / Super Admin can also toggle to personal attendance using ?view=my_attendance.
    show_personal = (not user_is_admin) or (requested_view == "my_attendance")

    if show_personal:
        employee = getattr(request.user, "employee_profile", None)
        if not employee:
            if user_is_admin:
                return redirect("attendance_report")
            messages.error(
                request,
                "Your login account is not linked to an employee profile. Please contact HR.",
            )
            return redirect("dashboard")

        today = timezone.localdate()

        # Handle "Put Attendance" action with Wi-Fi & Geofence Verification
        if request.method == "POST":
            action = request.POST.get("action")
            now_time = timezone.now()

            if action == "put_attendance":
                mode = hospital_setting.verification_mode
                is_wifi_ok = is_on_hospital_wifi

                # Evaluate GPS if provided
                user_lat = request.POST.get("latitude")
                user_lon = request.POST.get("longitude")
                is_gps_ok = False
                distance = None
                dist_str = ""

                if user_lat and user_lon:
                    try:
                        u_lat = float(user_lat)
                        u_lon = float(user_lon)
                        distance = calculate_haversine_distance(
                            hospital_setting.latitude,
                            hospital_setting.longitude,
                            u_lat,
                            u_lon,
                        )
                        dist_str = (
                            f"{int(distance)}m"
                            if distance < 1000
                            else f"{distance/1000:.2f} km"
                        )
                        if distance <= hospital_setting.radius_meters:
                            is_gps_ok = True
                    except (ValueError, TypeError):
                        pass

                # Enforcement by Verification Mode
                verification_source = ""
                if mode == HospitalSetting.VerificationMode.WIFI_ONLY:
                    if not is_wifi_ok:
                        messages.error(
                            request,
                            f"Attendance Rejected: You must be connected to the '{hospital_setting.wifi_name}' network. "
                            f"(Your detected Network IP: {client_ip}). Please connect to Hospital Wi-Fi and retry.",
                        )
                        return redirect("attendance_report")
                    verification_source = f"Hospital Wi-Fi ({hospital_setting.wifi_name})"

                elif mode == HospitalSetting.VerificationMode.GPS_ONLY:
                    if not is_gps_ok:
                        if distance is not None:
                            messages.error(
                                request,
                                f"Location verification failed: You are currently {dist_str} away from {hospital_setting.hospital_name}. "
                                f"You must be within {hospital_setting.radius_meters}m to mark attendance.",
                            )
                        else:
                            messages.error(
                                request,
                                "Location verification failed: GPS coordinates are required to verify you are on hospital grounds. Please allow browser location access.",
                            )
                        return redirect("attendance_report")
                    verification_source = "GPS Geofence"

                elif mode == HospitalSetting.VerificationMode.BOTH:
                    if not is_wifi_ok:
                        messages.error(
                            request,
                            f"Attendance Rejected: You must be connected to the '{hospital_setting.wifi_name}' network (Detected IP: {client_ip}).",
                        )
                        return redirect("attendance_report")
                    if not is_gps_ok:
                        messages.error(
                            request,
                            f"Attendance Rejected: You must also be within {hospital_setting.radius_meters}m of hospital premises ({dist_str or 'GPS not detected'}).",
                        )
                        return redirect("attendance_report")
                    verification_source = f"Hospital Wi-Fi + GPS"

                else:  # WIFI_OR_GPS (Either Wi-Fi OR GPS)
                    if not is_wifi_ok and not is_gps_ok:
                        messages.error(
                            request,
                            f"Attendance Rejected: You must either be connected to '{hospital_setting.wifi_name}' or physically within {hospital_setting.radius_meters}m of hospital premises. (Detected IP: {client_ip}).",
                        )
                        return redirect("attendance_report")
                    if is_wifi_ok and is_gps_ok:
                        verification_source = "Hospital Wi-Fi + GPS"
                    elif is_wifi_ok:
                        verification_source = f"Hospital Wi-Fi ({hospital_setting.wifi_name})"
                    else:
                        verification_source = "GPS Geofence"

                shift = (
                    employee.department.shifts.first()
                    if employee.department
                    else None
                )
                today_record, created = Attendance.objects.get_or_create(
                    employee=employee,
                    date=today,
                    defaults={
                        "shift": shift,
                        "check_in": now_time,
                        "latitude": float(user_lat) if user_lat else None,
                        "longitude": float(user_lon) if user_lon else None,
                        "ip_address": client_ip,
                        "verification_source": verification_source,
                    },
                )

                if not created and not today_record.check_in:
                    today_record.check_in = now_time
                    today_record.latitude = float(user_lat) if user_lat else None
                    today_record.longitude = float(user_lon) if user_lon else None
                    today_record.ip_address = client_ip
                    today_record.verification_source = verification_source
                    today_record.save(
                        update_fields=[
                            "check_in",
                            "latitude",
                            "longitude",
                            "ip_address",
                            "verification_source",
                        ]
                    )
                    messages.success(
                        request,
                        f"Attendance marked successfully for today at {timezone.localtime(now_time).strftime('%H:%M')} (Present) — Verified via {verification_source}.",
                    )
                elif not created:
                    messages.info(
                        request,
                        f"Attendance is already marked for today at {timezone.localtime(today_record.check_in).strftime('%H:%M')} (Present).",
                    )
                else:
                    messages.success(
                        request,
                        f"Attendance marked successfully for today at {timezone.localtime(now_time).strftime('%H:%M')} (Present) — Verified via {verification_source}.",
                    )

            return redirect("attendance_report")

        # Determine start date: from employee's date of joining up to today
        start_date = employee.joining_date or today
        first_record_date = (
            Attendance.objects.filter(employee=employee)
            .order_by("date")
            .values_list("date", flat=True)
            .first()
        )
        if first_record_date and first_record_date < start_date:
            start_date = first_record_date

        if start_date > today:
            start_date = today

        # Fetch all attendance records for employee from start_date to today
        records_by_date = {
            att.date: att
            for att in Attendance.objects.filter(
                employee=employee,
                date__gte=start_date,
                date__lte=today,
            ).select_related("shift")
        }

        # Build day-by-day history: from today down to date of joining
        attendance_history = []
        cur_date = today
        present_count = 0
        absent_count = 0

        while cur_date >= start_date:
            att = records_by_date.get(cur_date)
            is_present = att is not None and att.check_in is not None
            if is_present:
                present_count += 1
            else:
                absent_count += 1

            attendance_history.append({
                "date": cur_date,
                "day": cur_date.strftime("%A"),
                "status": "Present" if is_present else "Absent",
                "is_present": is_present,
                "check_in": att.check_in if att else None,
                "check_out": att.check_out if att else None,
                "verification_source": att.verification_source if att else None,
                "ip_address": att.ip_address if att else None,
            })
            cur_date -= timedelta(days=1)

        today_record = records_by_date.get(today)
        today_marked = (
            today_record is not None and today_record.check_in is not None
        )

        context = {
            "employee": employee,
            "attendance_history": attendance_history,
            "today_marked": today_marked,
            "today_record": today_record,
            "present_count": present_count,
            "absent_count": absent_count,
            "joining_date": employee.joining_date or start_date,
            "hospital_setting": hospital_setting,
            "client_ip": client_ip,
            "is_on_hospital_wifi": is_on_hospital_wifi,
            "is_hr_admin": user_is_admin,
        }
        return render(request, "hr/attendance_personal.html", context)

    # ---------------------------------------------------------
    # HR & SUPER ADMIN: Full Hospital-Wide Attendance Report
    # ---------------------------------------------------------
    today = timezone.localdate()
    date_value = request.GET.get(
        "date",
        str(today),
    )

    try:
        selected_date = date.fromisoformat(date_value)
    except (ValueError, TypeError):
        selected_date = today

    departments = Department.objects.all()

    selected_department = request.GET.get(
        "department",
        "",
    )

    query = request.GET.get("q", "").strip()

    attendance_records = Attendance.objects.filter(
        date=selected_date
    ).select_related(
        "employee",
        "employee__department",
        "shift",
    )

    if selected_department:
        attendance_records = attendance_records.filter(
            employee__department_id=selected_department
        )

    if query:
        attendance_records = attendance_records.filter(
            Q(employee__name__icontains=query)
            | Q(employee__employee_id__icontains=query)
        )

    dept_obj = (
        departments.filter(id=selected_department).first()
        if selected_department
        else None
    )
    department_name = dept_obj.name if dept_obj else "All Departments"

    # Handle Export requests for Hospital-Wide Report
    export_format = request.GET.get("export", "").lower()
    if export_format == "excel":
        return export_attendance_excel(
            attendance_records, selected_date, department_name
        )
    elif export_format == "pdf":
        return export_attendance_pdf(
            attendance_records, selected_date, department_name
        )

    context = {
        "attendance_records": attendance_records,
        "departments": departments,
        "selected_date": selected_date,
        "selected_department": selected_department,
        "query": query,

        "present_count": attendance_records.filter(
            check_in__isnull=False
        ).count(),

        "employees_count": attendance_records.values(
            "employee_id"
        ).distinct().count(),

        "export_excel_url": f"?date={selected_date.isoformat()}&department={selected_department}&q={query}&export=excel",
        "export_pdf_url": f"?date={selected_date.isoformat()}&department={selected_department}&q={query}&export=pdf",
        "has_employee_profile": bool(getattr(request.user, "employee_profile", None)),
        "hospital_setting": hospital_setting,
        "client_ip": client_ip,
        "is_on_hospital_wifi": is_on_hospital_wifi,
        "is_hr_admin": True,
    }

    return render(
        request,
        "hr/attendance_report.html",
        context,
    )


# =====================================================
# 7. LEAVE REQUEST LIST
# =====================================================

@login_required
def leave_list(request):
    user_is_admin = is_hr_or_admin(request.user)

    leave_requests = LeaveRequest.objects.select_related(
        "employee", "reviewed_by"
    ).order_by("-start_date")

    # Regular employees: ONLY see their own personal leave requests!
    if not user_is_admin:
        employee = getattr(request.user, "employee_profile", None)
        if employee:
            leave_requests = leave_requests.filter(employee=employee)
        else:
            leave_requests = leave_requests.filter(employee__user=request.user)

    status = request.GET.get("status", "")
    query = request.GET.get("q", "").strip()

    if status in ("PENDING", "APPROVED", "REJECTED"):
        leave_requests = leave_requests.filter(
            status=status
        )

    if query:
        leave_requests = leave_requests.filter(
            Q(employee__name__icontains=query)
            | Q(employee__employee_id__icontains=query)
            | Q(reason__icontains=query)
        )

    return render(
        request,
        "hr/leave_list.html",
        {
            "leave_requests": leave_requests,
            "selected_status": status,
            "query": query,
            "can_approve_leaves": user_is_admin and request.user.has_perm(
                "hr.approve_leave_requests"
            ),
            "is_hr_admin": user_is_admin,
        },
    )


# =====================================================
# 8. CREATE LEAVE REQUEST
# =====================================================

@login_required
def leave_create(request):
    employee = getattr(request.user, "employee_profile", None) or Employee.objects.filter(
        user=request.user
    ).first()

    if employee is None:
        messages.error(
            request,
            "Your account is not linked to an employee profile. "
            "Contact the HR administrator.",
        )
        return redirect("dashboard")

    form = LeaveRequestForm(request.POST or None)

    if request.method == "POST" and form.is_valid():
        leave = form.save(commit=False)
        leave.employee = employee
        leave.status = "PENDING"
        leave.save()

        messages.success(
            request,
            "Leave request submitted successfully. "
            "It is awaiting approval.",
        )

        return redirect("leave_list")

    return render(
        request,
        "hr/leave_form.html",
        {"form": form},
    )


# =====================================================
# 9. APPROVE OR REJECT LEAVE
# =====================================================

@login_required
@permission_required(
    "hr.approve_leave_requests",
    raise_exception=True,
)
def leave_review(request, pk):
    if request.method != "POST":
        return redirect("leave_list")

    leave = get_object_or_404(
        LeaveRequest.objects.select_related(
            "employee",
            "employee__user",
            "employee__manager",
        ),
        pk=pk,
    )

    # Never allow a person to approve their own leave.
    if leave.employee.user_id == request.user.id:
        messages.error(
            request,
            "You cannot approve or reject your own leave request.",
        )
        return redirect("leave_list")

    # Superusers with the approval permission can review any
    # employee's request, but still cannot review their own.
    # Other approvers must be the employee's assigned manager.
    if not request.user.is_superuser:
        is_manager = Employee.objects.filter(
            pk=leave.employee.manager_id,
            user=request.user,
        ).exists()

        if not is_manager:
            messages.error(
                request,
                "You are not authorized to review this employee's "
                "leave request. Only the assigned manager can do so.",
            )
            return redirect("leave_list")

    if leave.status != "PENDING":
        messages.warning(
            request,
            "This leave request has already been reviewed.",
        )
        return redirect("leave_list")

    decision = request.POST.get("decision")

    if decision not in ("APPROVED", "REJECTED"):
        messages.error(
            request,
            "Invalid leave decision.",
        )
        return redirect("leave_list")

    leave.status = decision
    leave.reviewed_by = request.user
    leave.reviewed_at = timezone.now()

    leave.save(
        update_fields=[
            "status",
            "reviewed_by",
            "reviewed_at",
        ]
    )

    messages.success(
        request,
        f"Leave request {decision.lower()} successfully.",
    )

    return redirect("leave_list")


# =====================================================
# 10. BIRTHDAY CALENDAR
# =====================================================

@login_required
@permission_required(
    "hr.view_birthday_calendar",
    raise_exception=True,
)
def birthday_calendar(request):
    today = timezone.localdate()

    month_value = request.GET.get(
        "month",
        today.strftime("%Y-%m"),
    )

    try:
        year, month = map(
            int,
            month_value.split("-"),
        )

        if not 1 <= month <= 12:
            raise ValueError

        if not 1 <= year <= 9999:
            raise ValueError

        selected_month = f"{year:04d}-{month:02d}"

    except (ValueError, TypeError):
        year = today.year
        month = today.month
        selected_month = today.strftime("%Y-%m")

    birthdays = Employee.objects.filter(
        status="ACTIVE",
        birthday__month=month,
    ).select_related(
        "department"
    ).order_by(
        "birthday__day",
        "name",
    )

    return render(
        request,
        "hr/birthday_calendar.html",
        {
            "birthdays": birthdays,
            "selected_month": selected_month,
            "month_title": date(
                year,
                month,
                1,
            ).strftime("%B %Y"),
        },
    )


# =====================================================
# 11. ROLE AND PERMISSION MANAGEMENT
# Superusers only
# =====================================================

ROLE_PERMISSION_OPTIONS = [
    ("view_dashboard", "View Dashboard"),
    ("view_employee_pages", "View Employees"),
    ("manage_employee_pages", "Add and Edit Employees"),
    ("view_attendance_pages", "View Attendance"),
    ("view_leave_pages", "View and Apply for Leave"),
    ("approve_leave_requests", "Approve or Reject Leave"),
    ("view_birthday_calendar", "View Birthday Calendar"),
    ("apply_leave_requests", "Apply for Leave"),
]


@login_required
@user_passes_test(lambda user: user.is_superuser)
def role_management(request):
    permission_labels = dict(
        ROLE_PERMISSION_OPTIONS
    )

    permission_options = [
        {
            "codename": codename,
            "label": label,
        }
        for codename, label in ROLE_PERMISSION_OPTIONS
    ]

    valid_codenames = set(
        permission_labels.keys()
    )

    selected_group = None
    selected_permissions = []
    role_name = ""

    group_id = (
        request.POST.get("group_id")
        or request.GET.get("group_id")
    )

    if group_id:
        selected_group = get_object_or_404(
            Group,
            pk=group_id,
        )

        role_name = selected_group.name

        selected_permissions = list(
            selected_group.permissions.filter(
                content_type__app_label="hr"
            ).values_list(
                "codename",
                flat=True,
            )
        )

    if request.method == "POST":
        role_name = request.POST.get(
            "role_name",
            "",
        ).strip()

        submitted_permissions = request.POST.getlist(
            "permissions"
        )

        selected_permissions = submitted_permissions

        duplicate_roles = Group.objects.filter(
            name__iexact=role_name
        )

        if selected_group:
            duplicate_roles = duplicate_roles.exclude(
                pk=selected_group.pk
            )

        if not role_name:
            messages.error(
                request,
                "Role name is required.",
            )

        elif duplicate_roles.exists():
            messages.error(
                request,
                "A role with this name already exists.",
            )

        elif not set(submitted_permissions).issubset(
            valid_codenames
        ):
            messages.error(
                request,
                "Invalid permission selected.",
            )

        else:
            permissions = Permission.objects.filter(
                content_type__app_label="hr",
                codename__in=submitted_permissions,
            )

            found_codenames = set(
                permissions.values_list(
                    "codename",
                    flat=True,
                )
            )

            missing_codenames = (
                set(submitted_permissions)
                - found_codenames
            )

            if missing_codenames:
                messages.error(
                    request,
                    "Some permissions are missing from the database: "
                    + ", ".join(
                        sorted(missing_codenames)
                    )
                    + ". Run migrations and check the custom permissions.",
                )

            else:
                with transaction.atomic():
                    role = selected_group or Group()
                    role.name = role_name
                    role.save()
                    role.permissions.set(permissions)

                messages.success(
                    request,
                    f"Role '{role.name}' saved successfully.",
                )

                return redirect("role_management")

    groups = Group.objects.prefetch_related(
        "permissions__content_type",
        "user_set",
    ).order_by("name")

    for group in groups:
        hr_permissions = [
            permission
            for permission in group.permissions.all()
            if permission.content_type.app_label == "hr"
        ]

        group.hr_permission_labels = [
            permission_labels.get(
                permission.codename,
                permission.name,
            )
            for permission in hr_permissions
        ]

        group.hr_permission_count = len(
            hr_permissions
        )

        group.assigned_users = list(
            group.user_set.all()
        )

        group.assigned_user_count = len(
            group.assigned_users
        )

    context = {
        "groups": groups,
        "selected_group": selected_group,
        "selected_permissions": selected_permissions,
        "permission_options": permission_options,
        "permission_options_codenames": sorted(
            valid_codenames
        ),
        "role_name": role_name,
    }

    return render(
        request,
        "hr/role_management.html",
        context,
    )